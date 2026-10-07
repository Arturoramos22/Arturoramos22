"""Chequeo aproximado de desborde de texto en un .pptx (sin renderizador disponible)."""
import sys, math
from pptx import Presentation
from pptx.util import Emu
EMU=914400
def est(text, size_pt, width_in, bold=False):
    # ancho medio de carácter Calibri ~0.50 em; Cambria ~0.55
    cw = size_pt/72*0.50*(1.06 if bold else 1)
    cpl = max(1, int(width_in / cw))
    lines=0
    for para in text.split("\n"):
        words=para.split(" "); cur=0; l=1
        for w in words:
            if cur+len(w) > cpl and cur>0: l+=1; cur=len(w)+1
            else: cur+=len(w)+1
        lines+=l
    return lines
prs=Presentation(sys.argv[1])
sw=prs.slide_width/EMU; sh=prs.slide_height/EMU
issues=0
for i,slide in enumerate(prs.slides,1):
    for shp in slide.shapes:
        if not shp.has_text_frame: continue
        tf=shp.text_frame
        if not tf.text.strip(): continue
        w=shp.width/EMU; h=shp.height/EMU; x=shp.left/EMU; y=shp.top/EMU
        if x<0 or y<0 or x+w>sw+0.01 or y+h>sh+0.01:
            print(f"S{i} FUERA DE LIENZO: '{tf.text[:40]}' x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}"); issues+=1
        need=0.0
        for p in tf.paragraphs:
            txt="".join(r.text for r in p.runs)
            if not txt.strip(): 
                need+= (p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 14)*1.2/72; continue
            size=None; bold=False
            for r in p.runs:
                if r.font.size: size=max(size or 0, r.font.size.pt)
                if r.font.bold: bold=True
            size=size or 14
            lines=est(txt,size,max(0.1,w-0.1),bold)
            sa=(p.space_after.pt if p.space_after else 0)
            need+=lines*size*1.2/72 + sa/72
        if need>h*1.02+0.02:
            print(f"S{i} POSIBLE DESBORDE: '{tf.text[:45]}...' necesita {need:.2f}in, caja {h:.2f}in (w={w:.2f})"); issues+=1
print("issues:",issues)
