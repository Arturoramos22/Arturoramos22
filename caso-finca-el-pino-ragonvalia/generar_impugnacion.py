# -*- coding: utf-8 -*-
"""Genera IMPUGNACION_TUTELA_2026-00107_v2.docx con imágenes incrustadas y notas al pie reales."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from xml.sax.saxutils import escape

S = "/tmp/claude-0/-home-user-Arturoramos22/47d246c2-9793-50f1-a1db-41f2c65e90c2/scratchpad"
OUT = f"{S}/IMPUGNACION_TUTELA_2026-00107_v2.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3); sec.right_margin = Cm(2.5)

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

FOOTNOTES = []  # list of strings


def add_footnote_ref(paragraph, text):
    """Inserta una referencia a nota al pie real en el párrafo y guarda el texto."""
    FOOTNOTES.append(text)
    fid = len(FOOTNOTES)
    run = paragraph.add_run()
    rPr = OxmlElement("w:rPr")
    va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript"); rPr.append(va)
    run._r.append(rPr)
    ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), str(fid))
    run._r.append(ref)
    return run


def build_footnotes_part(document):
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    parts = [f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
             f'<w:footnotes xmlns:w="{W}">',
             '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>',
             '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>']
    for i, txt in enumerate(FOOTNOTES, start=1):
        parts.append(
            f'<w:footnote w:id="{i}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/><w:jc w:val="both"/></w:pPr>'
            f'<w:r><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:vertAlign w:val="superscript"/><w:sz w:val="16"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:sz w:val="16"/></w:rPr><w:t xml:space="preserve"> {escape(txt)}</w:t></w:r></w:p></w:footnote>')
    parts.append('</w:footnotes>')
    xml = "\n".join(parts).encode("utf-8")
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                xml, document.part.package)
    document.part.relate_to(part, RT.FOOTNOTES)


def P(text="", bold=False, italic=False, align="justify", size=None, space_after=6, indent=0):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    return p


def RUNS(parts, align="justify", indent=0, space_after=6):
    """parts: lista de (texto, {'b':..,'i':..}) o ('FN', 'texto de la nota')."""
    p = P(align=align, indent=indent, space_after=space_after)
    for item in parts:
        if isinstance(item, tuple) and item[0] == "FN":
            add_footnote_ref(p, item[1]); continue
        text, fmt = (item, {}) if isinstance(item, str) else item
        r = p.add_run(text); r.bold = fmt.get("b", False); r.italic = fmt.get("i", False)
    return p


def H(text, level=1):
    p = P(align="left" if level > 1 else "center", space_after=6)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5 if level == 1 else 11)
    return p


def QUOTE(text, size=10):
    p = P(align="justify", indent=1.25, space_after=6)
    p.paragraph_format.right_indent = Cm(0.75)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(size)
    return p


def IMG(path, caption, width_cm=15.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9)
    return p


# ------------------------------------------------------------------ ENCABEZADO
P("Señor", align="left", space_after=0)
P("JUEZ PROMISCUO MUNICIPAL DE RAGONVALIA (NORTE DE SANTANDER)", bold=True, align="left", space_after=0)
P("Para ante el superior funcional", align="left", space_after=0)
P("E. S. D.", align="left", space_after=12)

tbl = doc.add_table(rows=5, cols=2); tbl.alignment = WD_TABLE_ALIGNMENT.LEFT; tbl.autofit = False
tbl.columns[0].width = Cm(3.6); tbl.columns[1].width = Cm(12.4)
rows = [("REFERENCIA:", "IMPUGNACIÓN del fallo de tutela de primera instancia del 2 de septiembre de 2026"),
        ("RADICADO:", "54-599-40-89-001-2026-00107-00"),
        ("ACCIONANTE:", "FREDDY ENRIQUE CARRILLO RINCÓN"),
        ("ACCIONADA:", "ALCALDÍA MUNICIPAL DE RAGONVALIA"),
        ("VINCULADAS:", "Secretaría de Planeación e Infraestructura de Ragonvalia y otras")]
for i, (a, b) in enumerate(rows):
    ca, cb = tbl.rows[i].cells
    ca.width = Cm(3.6); cb.width = Cm(12.4)
    ca.paragraphs[0].add_run(a).bold = True
    cb.paragraphs[0].add_run(b)
    for c in (ca, cb):
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
P("", space_after=6)

RUNS([("FREDDY ENRIQUE CARRILLO RINCÓN", {"b": True}),
      ", mayor de edad, identificado con la cédula de ciudadanía No. 13.453.598 expedida en Cúcuta, actuando en causa propia, "
      "dentro del término del artículo 31 del Decreto 2591 de 1991, presento ",
      ("IMPUGNACIÓN", {"b": True}),
      " contra la sentencia del 2 de septiembre de 2026, que declaró improcedente la acción de tutela de la referencia, "
      "para que el superior la revoque y, en su lugar, ampare mis derechos fundamentales, con fundamento en lo siguiente."])

# ------------------------------------------------------------------ I. OPORTUNIDAD
H("I. OPORTUNIDAD DE LA IMPUGNACIÓN")
RUNS(["La sentencia me fue notificada el [●] de septiembre de 2026 por correo electrónico. El presente escrito se radica dentro de los tres (3) días "
      "siguientes, conforme al artículo 31 del Decreto 2591 de 1991. Solicito que, de conformidad con el artículo 32 ibídem, se remita el "
      "expediente al superior funcional dentro de los dos (2) días siguientes."])

# ------------------------------------------------------------------ II. SÍNTESIS
H("II. LO QUE DECIDIÓ EL FALLO Y POR QUÉ DEBE REVOCARSE")
RUNS(["El fallo declaró improcedente la acción por incumplimiento del requisito de subsidiariedad sobre dos premisas: ",
      ("(i)", {"b": True}), " que la controversia versa sobre el Decreto 018 de 2023 y que, para controvertirlo, dispongo del medio de control de "
      "nulidad y restablecimiento del derecho con su régimen de medidas cautelares; y ",
      ("(ii)", {"b": True}), " que no existe perjuicio irremediable porque la Secretaría de Planeación e Infraestructura informó que "
      "\"actualmente no se adelantan obras\" sobre el predio y que la vía es \"un corredor vial preexistente\" incorporado al inventario vial municipal."])
RUNS(["Ambas premisas son equivocadas. La primera altera el objeto de la tutela: no pedí que el juez constitucional anulara ningún acto administrativo, "
      "sino que detuviera una intervención material ejecutada sin acto, sin procedimiento y sin título, y que garantizara el debido proceso "
      "administrativo que la propia entidad confesó haber omitido. La segunda es el resultado de una valoración probatoria que tuvo por cierta, sin "
      "prueba alguna, la versión de la entidad vinculada, y que pasó por alto la confesión de esa misma entidad sobre el ingreso de su maquinaria, "
      "el silencio de la Alcaldía accionada y las pruebas que acompañan la demanda. Explico cada cargo."])

# ------------------------------------------------------------------ III. CARGOS
H("III. RAZONES DE LA IMPUGNACIÓN")

# ---- CARGO 1
H("Primer cargo. El fallo cambió el objeto de la tutela: lo que se ataca es una vía de hecho y una omisión, no la legalidad de un acto administrativo", 2)
RUNS(["El fallo afirma que mis pretensiones \"se encuentran encaminadas, en esencia, a cuestionar\" el Decreto 018 de 2023 y que por ello el juez natural "
      "es el contencioso administrativo. No es así. Las pretensiones segunda, tercera y cuarta de la demanda pidieron, en su orden, la suspensión de la "
      "intervención material sobre el predio, la expedición y notificación del acto particular que defina la afectación y la garantía del debido proceso, "
      "la defensa y la contradicción. Ninguna de ellas pide que se declare la nulidad del decreto ni de oficio alguno."])
RUNS(["La diferencia no es de forma. Contra un acto administrativo existe, en efecto, el medio de control de nulidad y restablecimiento del derecho. "
      "Pero lo que aquí ocurre es distinto: la Administración interviene materialmente un inmueble privado ",
      ("sin acto administrativo que lo ordene, sin procedimiento previo y sin título", {"b": True}),
      ". Así lo confesó por escrito. En el oficio PLA 000298 del 8 de abril de 2026, transcrito íntegramente en la sentencia de tutela del 23 de abril de "
      "2026 (rad. 2026-00031), la Secretaría de Planeación respondió que sobre la declaratoria de la vía, la indemnización y las actas de visita "
      "\"no se encontró en el archivo municipal información al respecto\":"])
IMG(f"{S}/crops/t1_pla298.png",
    "Imagen 1. Oficio PLA 000298 del 8 de abril de 2026, transcrito en la sentencia de tutela del 23 de abril de 2026 (rad. 2026-00031), pág. 13.", 13.5)
RUNS(["Y en el oficio PLA 000355 del 25 de abril de 2026, expedido en cumplimiento de ese fallo, reconoció que \"no existen soportes respecto de los "
      "posibles actos en los que se produjo cesión voluntaria, expropiación por vía judicial o administrativa, constitución de servidumbre legal o de "
      "tránsito o cualquier otro mecanismo\", y que \"no se encontró documento que dé cuenta de la existencia de indemnización o trámite compensatorio "
      "por afectación de su predio\". El Instituto Nacional de Vías, por su parte, certificó que no existe red terciaria a su cargo en Ragonvalia:"])
IMG(f"{S}/crops/invias_no_red.png",
    "Imagen 2. Oficio INVIAS 2026S-VBOG-010160 del 2 de marzo de 2026, pág. 2, aportado por el INVIAS en la tutela rad. 2026-00031.", 15)
RUNS(["El único \"título\" que la Administración ha invocado alguna vez es una supuesta \"servidumbre de hecho\" por la antigüedad del camino, alegada "
      "al contestar la primera tutela:"])
IMG(f"{S}/crops/t1_servidumbre.png",
    "Imagen 3. Sentencia de tutela del 23 de abril de 2026 (rad. 2026-00031), acápite 4.1, síntesis de la contestación de la Secretaría de Planeación.", 15)
RUNS(["Esa figura no existe en el derecho colombiano. La servidumbre de tránsito es discontinua (art. 881 del Código Civil) y el artículo 939 del mismo "
      "código dispone que las servidumbres discontinuas \"sólo pueden adquirirse por medio de un título; ni aun el goce inmemorial bastará para "
      "constituirlas\". El paso que desde hace unos años permito por mera tolerancia a un vecino y, a partir de él, a algunos miembros de la comunidad, "
      "no convierte un camino interno de mi finca en vía pública."])
RUNS(["Frente a una intervención de hecho no hay acto que demandar. El medio de control de nulidad y restablecimiento exige un acto administrativo "
      "particular, que aquí no existe respecto de las obras. El de reparación directa opera después de consumado el daño y sólo lo indemniza. Ninguno de "
      "los dos impide, hoy, que la maquinaria municipal siga entrando a mi predio. La Corte Constitucional ha sido categórica: el Estado no puede ocupar "
      "bienes privados por vías de hecho; sólo puede hacerlo por motivos de utilidad pública, mediante expropiación con indemnización previa, y la "
      "vulneración de ese principio compromete su responsabilidad",
      ("FN", "Corte Constitucional, sentencia T-696 de 2010 (M.P. Juan Carlos Henao Pérez), proferida a propósito de la ocupación de predios privados para vías "
             "de la red terciaria sin expropiación previa: \"sólo por motivos de utilidad pública o de interés social se autoriza al Estado para acudir a la "
             "figura de la expropiación, para la cual debe mediar sentencia judicial e indemnización previa (…) nada justifica la ocupación de bienes de propiedad "
             "privada por parte del Estado mediante vías de hecho\"."),
      ". La Ley 388 de 1997 (arts. 58 a 62) y la Ley 1228 de 2008 (art. 3, parágrafo 1) señalan el cauce: si el Municipio necesita el corredor para una vía "
      "terciaria, debe adquirir la franja y su faja de retiro por enajenación voluntaria o expropiación. El interés general de la vereda Agualinda, que no "
      "discuto, se satisface por ese cauce y no mediante el sacrificio de un solo administrado sin procedimiento ni indemnización."])
RUNS(["Por eso el fallo dejó sin respuesta la pretensión central de la tutela. La garantía del debido proceso administrativo (art. 29 C.P.) exige que, antes "
      "de afectar un inmueble determinado, la Administración adelante una actuación con citación del titular (arts. 35 y 37 del CPACA), adopte una "
      "decisión motivada y la notifique en debida forma. Nada de eso ha ocurrido. El juez constitucional es competente para ordenar que ese procedimiento se "
      "adelante, porque contra una ",
      ("omisión", {"b": True}), " no hay acto que demandar ante el contencioso. El fallo no estudió esta pretensión; se limitó a remitirme a un medio de "
      "control que no tiene objeto sobre el cual recaer."])

# ---- CARGO 2
H("Segundo cargo. La subsidiariedad se examinó en abstracto y no en concreto: el medio ordinario ya está en curso y no ofrece protección inmediata", 2)
RUNS(["El artículo 6 del Decreto 2591 de 1991 ordena que la existencia de otros medios de defensa \"será apreciada en concreto, en cuanto a su eficacia, "
      "atendiendo las circunstancias en que se encuentre el solicitante\". El fallo no hizo ese examen. Se limitó a enunciar, en abstracto, que el CPACA "
      "prevé medidas cautelares."])
RUNS(["Lejos de desconocer la jurisdicción contencioso administrativa, acudí a ella. El 26 de agosto de 2026 radiqué ante la Procuraduría General de la "
      "Nación la solicitud de conciliación extrajudicial, requisito de procedibilidad de la demanda de nulidad y restablecimiento del derecho contra el "
      "Municipio de Ragonvalia, con copia íntegra remitida ese mismo día a la Alcaldía; la Procuraduría la radicó bajo el número E-2026-483440 (acuse del "
      "27 de agosto de 2026, que se aporta). Esa demanda, cuyo texto acompaña la solicitud, pide la suspensión provisional de los oficios y del aparte "
      "cartográfico del EOT y una medida cautelar de urgencia para detener las obras."])
RUNS(["Pero ese cauce, que es el definitivo, ",
      ("no es eficaz para conjurar el daño en curso", {"b": True}),
      ". La medida cautelar sólo puede pedirse con la demanda; la demanda sólo puede presentarse cuando termine el trámite conciliatorio, que la ley "
      "permite extender hasta por tres meses; y aun radicada, la cautelar exige reparto, admisión y decisión. Entre tanto, la Administración, que sostiene "
      "que el camino es una vía pública que \"mantiene\" con su propia maquinaria, puede consumar la afectación de mi vivienda y de mi beneficiadero de café. "
      "Ese vacío temporal es exactamente lo que el amparo transitorio del artículo 8 del Decreto 2591 de 1991 está llamado a cubrir. La propia sentencia "
      "impugnada lo reconoce cuando transcribe que la tutela \"es procedente como (i) mecanismo transitorio para evitar un perjuicio irremediable\"."])
RUNS(["Debo señalar, además, que el 20 de agosto de 2026 este mismo Despacho negó la medida provisional de suspensión de las obras con el argumento de "
      "que lo pretendido era tema central de otro escenario procesal, y ahora declara improcedente la tutela porque ese otro escenario existe. El resultado "
      "práctico es que ninguna autoridad judicial ha examinado, hasta hoy, si la maquinaria municipal puede o no seguir entrando a mi finca."])

# ---- CARGO 3
H("Tercer cargo. El perjuicio irremediable se descartó con una valoración probatoria equivocada", 2)
RUNS(["Este es el punto decisivo. El fallo concluyó que no hay amenaza cierta ni inminente porque la Secretaría de Planeación informó que no se adelantan "
      "obras y que la vía es preexistente. Esa conclusión adolece de cinco defectos."])

RUNS([("1. Tuvo por cierta la versión de la vinculada sin prueba alguna.", {"b": True}),
      " La Secretaría acompañó un \"archivo fotográfico de mantenimiento de vía terciaria\" y un pantallazo del SINC (pdf 008, folios 50 y 52, según la "
      "relación probatoria del fallo). El primero, lejos de desvirtuar mis hechos, los confirma: documenta la operación de la maquinaria municipal sobre la vía. "
      "No aportó bitácora de maquinaria, orden de trabajo, contrato ni prueba alguna de que el tramo que atraviesa mi predio hubiera quedado por fuera de "
      "esa intervención. El fallo, sin embargo, transcribe su informe y lo convierte en hecho probado:"])
IMG(f"{S}/crops/fallo_p15_secretaria.png", "Imagen 4. Sentencia impugnada, pág. 15.", 15)
RUNS(["Frente a esa afirmación estaban mis hechos cuarto, décimo y decimocuarto, el registro fotográfico del expediente (folios 13 y 14 del pdf 003) y el "
      "video del 19 de agosto de 2026 que ahora aporto. El fallo no los valoró ni explicó por qué la palabra de la entidad pesa más que la prueba del "
      "accionante. En materia de tutela, cuando la versión de la autoridad y la del accionante se contradicen sobre un hecho determinante, el juez debe "
      "decretar las pruebas necesarias (arts. 19, 21 y 22 del Decreto 2591 de 1991), lo que aquí no ocurrió pese a que una simple inspección de la "
      "Inspección de Policía habría bastado."])

RUNS([("2. La propia Secretaría confesó el ingreso de la maquinaria municipal.", {"b": True}),
      " El informe reconoce que en junio de 2026 se realizaron \"actividades rutinarias de mantenimiento vial en algunos sectores, mediante el uso de "
      "maquinaria de propiedad municipal\", y sólo niega intervenciones \"en el tramo localizado en inmediaciones del predio\":"])
IMG(f"{S}/crops/fallo_p7_maquinaria.png", "Imagen 5. Sentencia impugnada, pág. 7, síntesis del informe de la Secretaría de Planeación e Infraestructura.", 15)
RUNS(["La distinción es artificiosa. La \"Vía 30\" que el Municipio reclama como suya es, en ese sector, el camino que atraviesa mi finca y pasa entre mi "
      "vivienda, los corrales y el beneficiadero; no hay \"inmediaciones\" del predio distintas del predio mismo. Mantener esa vía con maquinaria municipal "
      "es, precisamente, la intervención material que denuncio. La entidad no niega que la maquinaria opere sobre el corredor: niega que eso sea una obra. "
      "Ese es el punto en discusión, no una razón para no discutirlo."])
IMG(f"{S}/img/conc_p09_11_848x480.png",
    "Imagen 6. Camino interno del predio “El Pino”, graficado por el Municipio como “Vía 30”: discurre entre los corrales (izq.) y el beneficiadero (der.). "
    "Fotograma del video del 19 de agosto de 2026 que se aporta con esta impugnación.", 12.5)
IMG(f"{S}/img/conc_p10_14_848x480.png",
    "Imagen 7. Vivienda y construcciones del predio atravesadas por el camino. Fotograma del mismo video.", 12.5)

RUNS([("3. La Alcaldía accionada guardó silencio.", {"b": True}),
      " La constancia secretarial del 27 de agosto de 2026 lo registra: la Alcaldía Municipal de Ragonvalia, única accionada, y la Inspección de Policía, "
      "no respondieron:"])
IMG(f"{S}/crops/fallo_p7_silencio.png", "Imagen 8. Sentencia impugnada, pág. 7, constancia de las respuestas recibidas.", 15)
RUNS(["El artículo 20 del Decreto 2591 de 1991 dispone que, si el informe no es rendido dentro del plazo, \"se tendrán por ciertos los hechos\". La "
      "Alcaldía, que es quien ordena y ejecuta las obras con su maquinaria, no controvirtió los hechos de la demanda. El fallo no aplicó la presunción de "
      "veracidad; hizo lo contrario: suplió el silencio de la accionada con el informe de una dependencia vinculada, y con ese informe desvirtuó mis hechos."])

RUNS([("4. La \"preexistencia\" de la vía es una afirmación sin soporte, contradicha por los propios documentos del Municipio.", {"b": True}),
      " El informe dice que la vía hace parte del inventario vial \"desde 2017\"; el oficio PLA 000355 dice que está registrada en el SINC \"desde el 2015\"; "
      "el oficio PLA 000298 dice que no se encontró la declaratoria; y el propio Decreto 018 de 2023, en su Programa de Ejecución, presupuesta para el "
      "mediano plazo 2028-2031 el proyecto \"Inventario y actualización de la red vial terciaria municipal e incorporación al mapa base\"",
      ("FN", "Decreto Municipal 018 del 20 de junio de 2023, Programa de Ejecución, pág. 161 de 183, línea estratégica \"Fortalecimiento del ordenamiento "
             "territorial y desarrollo urbano\", proyecto por $250.000.000 para el período 2028-2031."),
      ", es decir, reconoce que el inventario vial formal no existe todavía. Lo único que existe es un levantamiento topográfico contratado en noviembre "
      "de 2017 y un plano de consultoría de 2019, elaborados sin citarme y sin acto administrativo que los adopte. Un camino privado no se vuelve público "
      "porque un contratista lo dibuje en un plano."])

RUNS([("5. Los cuatro elementos del perjuicio irremediable están acreditados.", {"b": True}),
      " ", ("Inminencia:", {"i": True}), " la Administración sostiene que el camino es vía pública, que es el único acceso a la parte baja de la vereda Agualinda "
      "y que lo mantiene con su maquinaria; no es una posibilidad remota sino una posición institucional en ejecución. ",
      ("Gravedad:", {"i": True}), " el EOT y la Ley 1228 de 2008 asignan a las vías terciarias una faja de retiro de 30 metros, 15 a cada lado del eje; mi "
      "vivienda, las viviendas de los trabajadores y las marquesinas de secado del café están a pocos metros del eje (imágenes 6 y 7), de modo que la "
      "consolidación de la vía con esas especificaciones implica su demolición. ",
      ("Urgencia:", {"i": True}), " el daño ya está ocurriendo: el portón de acceso ha tenido que repararse al menos tres veces, se han hurtado 40 kilos de "
      "café seco, motores, herramientas y aves de corral, y han muerto cinco terneros y dos vacas por la apertura forzada del predio. ",
      ("Impostergabilidad:", {"i": True}), " una vez demolidas las edificaciones, ninguna sentencia contenciosa las devuelve; la indemnización posterior "
      "no restituye la vivienda ni la actividad de la que vive mi familia."])
RUNS(["A lo anterior se suma mi condición de campesino que habita y trabaja el predio, sujeto de especial protección constitucional conforme al artículo 64 "
      "de la Constitución, reformado por el Acto Legislativo 01 de 2023, y la naturaleza de los derechos comprometidos: vivienda digna, trabajo y mínimo "
      "vital, cuya afectación no es meramente patrimonial. El fallo trató el caso como una disputa sobre la propiedad; es, antes que eso, una disputa sobre "
      "el techo y el sustento de una familia rural."])

# ---- CARGO 4
H("Cuarto cargo. Contradicción interna del fallo", 2)
RUNS(["Al examinar la inmediatez, el fallo tuvo por satisfecho el requisito porque \"la controversia planteada no recae sobre un hecho consumado o "
      "superado, sino sobre una circunstancia cuyos efectos presuntamente continúan produciéndose en el tiempo\" (pág. 11). Cuatro páginas después, al "
      "examinar la subsidiariedad, afirmó que \"no existen elementos objetivos que permitan concluir que actualmente se esté ejecutando una actuación "
      "administrativa o material\" (pág. 15):"])
IMG(f"{S}/crops/fallo_p15_inminencia.png", "Imagen 9. Sentencia impugnada, pág. 15.", 15)
RUNS(["No puede ser al mismo tiempo una afectación actual, para efectos de inmediatez, y una afectación inexistente, para efectos de subsidiariedad. La "
      "contradicción revela que la segunda conclusión no fue el producto de un análisis probatorio sino de la aceptación acrítica del informe de la vinculada."])

# ---- PRECISIÓN
H("Precisión sobre el alcance de lo que se pide", 2)
RUNS(["No pido que el juez de tutela declare la nulidad del Decreto 018 de 2023 ni de los oficios PLA 000298 y PLA 000355: eso lo decidirá la jurisdicción "
      "de lo contencioso administrativo, ante la cual ya inicié el trámite. Tampoco pido que se prohíba el paso a la comunidad. Pido que, mientras esa "
      "jurisdicción decide, la Administración no consume con maquinaria una afectación que ella misma admite no tener cómo justificar, y que defina mi "
      "situación jurídica a través del procedimiento que la ley exige, con mi participación."])

# ------------------------------------------------------------------ IV. PETICIONES
H("IV. PETICIONES")
RUNS([("PRIMERA. REVOCAR", {"b": True}), " la sentencia del 2 de septiembre de 2026 proferida por el Juzgado Promiscuo Municipal de Ragonvalia."])
RUNS([("SEGUNDA. AMPARAR", {"b": True}), ", como mecanismo transitorio para evitar un perjuicio irremediable (art. 8 del Decreto 2591 de 1991), mis derechos "
      "fundamentales al debido proceso, a la vivienda digna, al trabajo y al mínimo vital, y, en conexidad con ellos, a la propiedad privada."])
RUNS([("TERCERA. ORDENAR", {"b": True}), " a la Alcaldía Municipal de Ragonvalia y a su Secretaría de Planeación e Infraestructura que, a partir de la "
      "notificación del fallo, se abstengan de ejecutar obras, ingresar maquinaria o realizar cualquier intervención material sobre el tramo del camino que "
      "discurre por el interior del predio \"El Pino\" (M.I. 264-19086), directamente o por medio de contratistas, hasta que la Jurisdicción de lo Contencioso "
      "Administrativo decida sobre las medidas cautelares o el fondo del proceso cuyo requisito de procedibilidad se encuentra en trámite ante la "
      "Procuraduría General de la Nación (rad. E-2026-483440). Manifiesto que presentaré la demanda dentro del término del artículo 8 del Decreto 2591 de "
      "1991, contado desde la expedición de la constancia que agote el requisito de procedibilidad."])
RUNS([("CUARTA. ORDENAR", {"b": True}), " a la Alcaldía Municipal de Ragonvalia que, en el término de treinta (30) días, defina mediante acto administrativo "
      "motivado, expedido con citación y audiencia del suscrito y notificado personalmente, si pretende adquirir el corredor y su faja de retiro conforme a "
      "la Ley 388 de 1997 o excluir el tramo de su cartografía e inventario vial; y que, mientras esa decisión no se adopte y quede en firme, se abstenga "
      "de toda intervención sobre el predio."])
RUNS([("QUINTA. SUBSIDIARIAMENTE", {"b": True}), ", si el superior estima que la prueba del perjuicio irremediable es insuficiente, solicito que antes de "
      "decidir, en ejercicio de las facultades de los artículos 19 y 21 del Decreto 2591 de 1991, decrete: (i) inspección judicial al predio o, en su defecto, informe "
      "de la Inspección de Policía de Ragonvalia con registro fotográfico fechado del tramo; y (ii) que la Alcaldía aporte la bitácora de operación de la "
      "maquinaria municipal entre junio y agosto de 2026, con fechas, sectores intervenidos y órdenes de trabajo, así como el contrato u orden que soporte "
      "las actividades ejecutadas sobre el corredor."])

# ------------------------------------------------------------------ V. PRUEBAS
H("V. PRUEBAS QUE SE APORTAN CON LA IMPUGNACIÓN")
RUNS(["Además de las que obran en el expediente, aporto las siguientes, que el superior puede valorar dado el carácter informal y sumario del trámite:"])
for t in ["Video del predio y del camino del 19 de agosto de 2026 (archivo VID-20260819-WA0014.mp4), del cual se extrajeron las imágenes 6 y 7.",
          "Registro fotográfico fechado de las intervenciones sobre el camino y del estado del portón, cercas y corrales [●].",
          "Solicitud de conciliación extrajudicial radicada el 26 de agosto de 2026 ante la Procuraduría General de la Nación y acuse de recibo con radicado E-2026-483440 del 27 de agosto de 2026.",
          "Oficio PLA 000355 del 25 de abril de 2026 con su correo remisorio y anexos (Decreto 018 de 2023, Plano No. 05 y plancha del levantamiento topográfico de noviembre de 2017).",
          "Oficio INVIAS 2026S-VBOG-010160 del 2 de marzo de 2026.",
          "Sentencia de tutela del 23 de abril de 2026, rad. 54-599-40-89-001-2026-00031-00."]:
    p = doc.add_paragraph(style="List Number"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3); p.add_run(t)

# ------------------------------------------------------------------ VI. NOTIFICACIONES
H("VI. NOTIFICACIONES")
pnot = RUNS(["Recibiré notificaciones en el correo electrónico carrillojaimesfreddy@gmail.com, en el teléfono 313 335 5383 y en la finca \"El Pino\", vereda "
      "Sombrerito, municipio de Ragonvalia."]); pnot.paragraph_format.keep_with_next = True
pa = P("Atentamente,", align="left", space_after=36); pa.paragraph_format.space_before = Pt(18); pa.paragraph_format.keep_with_next = True
pl = P("______________________________", align="left", space_after=0); pl.paragraph_format.keep_with_next = True
pn = P("FREDDY ENRIQUE CARRILLO RINCÓN", bold=True, align="left", space_after=0); pn.paragraph_format.keep_with_next = True
pc = P("C.C. No. 13.453.598 de Cúcuta", align="left", space_after=0); pc.paragraph_format.keep_with_next = True
P("Accionante", align="left")

build_footnotes_part(doc)
doc.save(OUT)
print("saved", OUT, "footnotes:", len(FOOTNOTES))
